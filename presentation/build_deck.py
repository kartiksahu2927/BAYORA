"""Build the Bayora hackathon pitch deck.

Run with:
    PYTHONPATH=/tmp/omnirush/pptxlib/site python3 presentation/build_deck.py

The deck intentionally uses editable PowerPoint text and vector shapes for the
story, with only the supplied Bayora artwork used as full-bleed visual texture.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT.parent / "Bayora_HIH_Nomination_Deck.pptx"

W = 13.333
H = 7.5

# Palette: high-contrast midnight with disciplined red/blue adversarial cues.
BG = "070B13"
BG_2 = "0B1220"
PANEL = "111B2B"
PANEL_2 = "17263B"
GRID = "1E2D44"
WHITE = "F4F7FB"
MUTED = "A8B4C6"
DIM = "70819A"
CYAN = "66F4D0"
BLUE = "5D9CFF"
RED = "FF536E"
YELLOW = "F7C86A"
PURPLE = "9A8CFF"
BLACK = "03050A"

FONT = "Aptos"
DISPLAY = "Aptos Display"
MONO = "Aptos Mono"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_transparency(fill, value: int) -> None:
    try:
        fill.transparency = value
    except Exception:
        # Transparency is cosmetic; the deck should still build on older
        # python-pptx versions.
        pass


def add_shape(slide, kind, x, y, w, h, fill=BG, line=None, transparency=0, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else kind
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    set_transparency(shape.fill, transparency)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    return shape


def rect(slide, x, y, w, h, fill=BG, line=None, transparency=0, radius=False):
    return add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, fill, line, transparency, radius)


def circle(slide, x, y, d, fill=BG, line=None, transparency=0):
    return add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, d, d, fill, line, transparency)


def line(slide, x1, y1, x2, y2, color=GRID, width=1.0, transparency=0, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    set_transparency(shape.line, transparency)
    if dash:
        try:
            shape.line.dash_style = dash
        except Exception:
            pass
    return shape


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=14,
    color=WHITE,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    italic=False,
    tracking=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    if tracking is not None:
        try:
            run.font._element.set("spc", str(tracking))
        except Exception:
            pass
    return box


def rich_text(slide, runs, x, y, w, h, size=14, font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    """Add a single paragraph with runs like {text, color, bold, font}."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    for item in runs:
        run = p.add_run()
        run.text = item["text"]
        run.font.name = item.get("font", font)
        run.font.size = Pt(item.get("size", size))
        run.font.bold = item.get("bold", False)
        run.font.italic = item.get("italic", False)
        run.font.color.rgb = rgb(item.get("color", WHITE))
    return box


def add_full_background(slide, image=None, tint=0):
    rect(slide, 0, 0, W, H, BG)
    if image:
        slide.shapes.add_picture(str(ASSETS / image), 0, 0, width=Inches(W), height=Inches(H))
        if tint:
            rect(slide, 0, 0, W, H, BG, transparency=tint)


def add_grid(slide, x=0, y=0, w=W, h=H, step=0.5, transparency=78):
    xx = x
    while xx <= x + w:
        line(slide, xx, y, xx, y + h, GRID, 0.45, transparency)
        xx += step
    yy = y
    while yy <= y + h:
        line(slide, x, yy, x + w, yy, GRID, 0.45, transparency)
        yy += step


def pill(slide, label, x, y, w, color=CYAN, fill=PANEL, text_color=WHITE, size=8.5):
    rect(slide, x, y, w, 0.26, fill, color, 0, True)
    text(slide, label.upper(), x, y + 0.01, w, 0.22, size, text_color, True, MONO, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.01)


def section_label(slide, kicker, title, subtitle=None, x=0.72, y=0.62, width=8.8):
    text(slide, kicker.upper(), x, y, width, 0.22, 8.5, CYAN, True, MONO, tracking=120)
    text(slide, title, x, y + 0.3, width, 0.72, 27, WHITE, True, DISPLAY, tracking=-10)
    if subtitle:
        text(slide, subtitle, x, y + 1.08, width, 0.52, 11.5, MUTED, False, FONT)


def chrome(slide, n, label="BAYORA / SECURE AI TESTING"):
    line(slide, 0.72, 0.36, 12.62, 0.36, GRID, 0.8)
    text(slide, label, 0.72, 0.12, 4.2, 0.18, 7.2, DIM, True, MONO, tracking=100)
    text(slide, f"{n:02d}", 12.08, 0.1, 0.5, 0.22, 8, CYAN, True, MONO, PP_ALIGN.RIGHT)
    line(slide, 0.72, 7.12, 12.62, 7.12, GRID, 0.8)
    text(slide, "ATTACK SURFACE → POLICY → EVIDENCE", 0.72, 7.22, 4.6, 0.15, 6.8, DIM, True, MONO, tracking=85)
    text(slide, "HACK IN HILLS / PITCH DECK", 9.25, 7.22, 3.37, 0.15, 6.8, DIM, True, MONO, PP_ALIGN.RIGHT, tracking=70)


def small_check(slide, x, y, label, color=CYAN):
    circle(slide, x, y, 0.22, color)
    text(slide, "✓", x, y - 0.01, 0.22, 0.22, 9, BG, True, FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
    text(slide, label, x + 0.32, y - 0.01, 3.4, 0.24, 10.5, WHITE, False, FONT, valign=MSO_ANCHOR.MIDDLE)


def numbered_card(slide, number, label, body, x, y, w, h, accent=CYAN):
    rect(slide, x, y, w, h, PANEL, GRID, 0, True)
    rect(slide, x, y, 0.07, h, accent)
    text(slide, number, x + 0.22, y + 0.18, 0.44, 0.3, 10, accent, True, MONO)
    text(slide, label.upper(), x + 0.22, y + 0.58, w - 0.42, 0.24, 10, WHITE, True, MONO, tracking=30)
    text(slide, body, x + 0.22, y + 0.91, w - 0.42, h - 1.04, 10.2, MUTED, False, FONT)


def add_arrow(slide, x, y, label="→", color=CYAN):
    text(slide, label, x, y, 0.34, 0.35, 20, color, True, FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)


def add_metric(slide, value, label, x, y, w, accent=CYAN):
    text(slide, value, x, y, w, 0.58, 28, accent, True, DISPLAY)
    text(slide, label.upper(), x, y + 0.62, w, 0.22, 7.6, MUTED, True, MONO, tracking=70)


def add_highlight(slide, x, y, w, h, color, transparency=80):
    rect(slide, x, y, w, h, color, transparency=transparency, radius=True)


def slide_01(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "hero.jpg", tint=28)
    # A deep reading plane keeps the pitch title legible over the artwork.
    rect(slide, 0, 0, 7.02, H, BG, transparency=7)
    rect(slide, 6.5, 0, 0.52, H, BG, transparency=44)
    add_grid(slide, 0.72, 0.75, 5.4, 5.75, 0.52, 88)
    line(slide, 0.72, 1.02, 1.82, 1.02, CYAN, 2.3)
    text(slide, "BAYORA", 0.72, 1.27, 5.8, 0.78, 44, WHITE, True, DISPLAY, tracking=-35)
    text(slide, "ATTACK THE MODEL.\nPROVE THE BOUNDARY.", 0.74, 2.22, 5.6, 1.5, 25, WHITE, True, DISPLAY, tracking=-15)
    # Color accents make the red-team / blue-team tension readable at a glance.
    rect(slide, 0.74, 3.92, 0.78, 0.08, RED)
    rect(slide, 1.62, 3.92, 0.78, 0.08, BLUE)
    text(
        slide,
        "Secure adversarial AI testing infrastructure for repeatable,\n evidence-backed security evaluation.",
        0.74,
        4.28,
        5.15,
        0.72,
        13,
        MUTED,
        False,
        FONT,
    )
    pill(slide, "RED TEAM", 0.74, 5.45, 1.1, RED, "2A121B")
    pill(slide, "BLUE TEAM", 1.98, 5.45, 1.18, BLUE, "101D39")
    pill(slide, "AUDIT", 3.28, 5.45, 0.88, CYAN, "0D2826")
    text(slide, "FROM PROMPT  →  POLICY  →  ISOLATED WORKER  →  AUDIT", 0.75, 6.2, 5.85, 0.24, 8.2, WHITE, True, MONO, tracking=32)
    text(slide, "HACK IN HILLS / HACKATHON PITCH", 0.74, 6.8, 4.8, 0.2, 7.4, DIM, True, MONO, tracking=92)
    text(slide, "01", 12.0, 0.45, 0.6, 0.24, 8, CYAN, True, MONO, PP_ALIGN.RIGHT)
    return slide


def slide_02(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "quiet.jpg", tint=10)
    chrome(slide, 2, "BAYORA / THE GAP")
    section_label(
        slide,
        "THE PROBLEM",
        "AI doesn’t fail in a lab.\nIt fails at the boundary.",
        "A crafted prompt, a poisoned document, or the wrong session can turn a useful model into a security incident.",
        0.72,
        0.68,
        6.6,
    )

    # Threat chain on the right.
    rect(slide, 8.35, 0.86, 4.28, 4.72, PANEL, GRID, 0, True)
    text(slide, "THE ATTACK SURFACE", 8.7, 1.18, 3.2, 0.22, 8.5, RED, True, MONO, tracking=90)
    line(slide, 8.7, 1.54, 12.3, 1.54, GRID, 0.8)
    chain = [
        ("01", "UNTRUSTED INPUT", "Prompt injection / poisoned content", RED),
        ("02", "MODEL CONTEXT", "Hidden instructions + sensitive state", YELLOW),
        ("03", "UNSAFE BEHAVIOR", "Leak, overreach, or unreviewable action", RED),
    ]
    yy = 1.86
    for num, label, body, accent in chain:
        circle(slide, 8.72, yy, 0.42, accent, None, 5)
        text(slide, num, 8.72, yy + 0.01, 0.42, 0.4, 8, BG, True, MONO, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
        text(slide, label, 9.35, yy - 0.01, 2.86, 0.22, 10.3, WHITE, True, MONO, tracking=28)
        text(slide, body, 9.35, yy + 0.28, 2.92, 0.34, 9.5, MUTED, False, FONT)
        if yy < 4:
            line(slide, 8.93, yy + 0.49, 8.93, yy + 0.76, accent, 1.2, 25)
        yy += 1.14
    pill(slide, "NO CONTROL PLANE", 8.7, 5.05, 1.72, RED, "2A121B")
    text(slide, "Teams get a screenshot — not a defensible trail.", 10.58, 5.05, 1.65, 0.4, 9.7, WHITE, True, FONT)

    # The real gap, summarized as three high-signal cards.
    text(slide, "THE GAP", 0.74, 3.9, 1.2, 0.2, 8.3, CYAN, True, MONO, tracking=100)
    numbered_card(slide, "01", "Hard to reproduce", "The same prompt can produce a different security story tomorrow.", 0.72, 4.28, 2.25, 1.58, CYAN)
    numbered_card(slide, "02", "Hard to scope", "Red, blue, model, tenant, and session boundaries blur together.", 3.12, 4.28, 2.25, 1.58, BLUE)
    numbered_card(slide, "03", "Hard to prove", "A response without provenance is not evidence a reviewer can trust.", 5.52, 4.28, 2.25, 1.58, RED)
    text(slide, "THREAT FRAMING: OWASP GENAI  •  MITRE ATLAS  •  RESPONSIBLE AI", 0.75, 6.32, 6.9, 0.18, 7.2, DIM, True, MONO, tracking=45)
    return slide


def slide_03(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "quiet.jpg", tint=12)
    chrome(slide, 3, "BAYORA / THE SOLUTION")
    section_label(
        slide,
        "THE SOLUTION",
        "Bayora makes every test traceable.",
        "One controlled loop turns adversarial input into a scoped decision, an isolated run, and evidence a reviewer can replay.",
        0.72,
        0.68,
        9.2,
    )

    # Contrast strip: the old workflow versus Bayora.
    rect(slide, 0.72, 2.08, 12.0, 0.82, PANEL, GRID, 0, True)
    text(slide, "BEFORE", 0.98, 2.23, 0.9, 0.22, 8.2, DIM, True, MONO, tracking=70)
    text(slide, "PROMPT  →  RESPONSE", 1.92, 2.19, 2.65, 0.3, 14, MUTED, True, DISPLAY)
    line(slide, 4.82, 2.49, 5.05, 2.49, GRID, 2.0)
    text(slide, "AFTER", 5.35, 2.23, 0.9, 0.22, 8.2, CYAN, True, MONO, tracking=70)
    text(slide, "PROBE  →  POLICY  →  WORKER  →  EVIDENCE  →  REVIEW", 6.27, 2.19, 5.95, 0.3, 13, WHITE, True, DISPLAY)

    text(slide, "THE BAYORA LOOP", 0.75, 3.37, 2.0, 0.22, 8.5, CYAN, True, MONO, tracking=90)
    steps = [
        ("01", "DEFINE", "Create an adversarial probe", RED),
        ("02", "SCOPE", "Resolve tenant + session policy", BLUE),
        ("03", "ISOLATE", "Run a bounded model worker", CYAN),
        ("04", "COLLECT", "Normalize output + provenance", PURPLE),
        ("05", "VERIFY", "Review the signed audit trail", YELLOW),
    ]
    x = 0.72
    for i, (num, label, body, accent) in enumerate(steps):
        numbered_card(slide, num, label, body, x, 3.78, 2.16, 1.74, accent)
        if i < len(steps) - 1:
            add_arrow(slide, x + 2.16, 4.4, "→", accent)
        x += 2.42

    # A small evidence artifact gives the slide a tangible payoff.
    rect(slide, 0.72, 5.98, 12.0, 0.7, "0D1C2C", CYAN, 0, True)
    circle(slide, 1.02, 6.2, 0.28, CYAN)
    text(slide, "✓", 1.02, 6.18, 0.28, 0.25, 10, BG, True, FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
    text(slide, "BOUNDARY TEST PASSED", 1.48, 6.14, 2.2, 0.22, 9.3, CYAN, True, MONO, tracking=40)
    text(slide, "session: 7f3c…b91a", 4.1, 6.16, 1.65, 0.22, 9.2, WHITE, False, MONO)
    text(slide, "decision: ALLOW", 6.06, 6.16, 1.55, 0.22, 9.2, WHITE, False, MONO)
    text(slide, "audit: HMAC-LINKED", 8.2, 6.16, 1.9, 0.22, 9.2, WHITE, False, MONO)
    text(slide, "reviewable", 11.05, 6.16, 1.25, 0.22, 9.2, CYAN, True, MONO, PP_ALIGN.RIGHT)
    return slide


def slide_04(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "arena.jpg", tint=76)
    chrome(slide, 4, "BAYORA / ARCHITECTURE")
    section_label(
        slide,
        "THE ARCHITECTURE",
        "Four trust zones. One controlled path.",
        "The API is the only component that receives client requests. Everything else is scoped, bounded, and observable.",
        0.72,
        0.64,
        8.5,
    )

    # Four zones keep the architecture legible in a live pitch.
    zones = [
        ("01", "CONTROL PLANE", "Browser\nAPI • Auth", CYAN, 0.72),
        ("02", "POLICY + PROVENANCE", "Tenant + session\nDefault deny • HMAC", BLUE, 3.72),
        ("03", "EVALUATION PLANE", "Isolated worker\nProvider adapter", RED, 6.72),
        ("04", "PROTECTED RESOURCES", "Other tenants\nHost • private data", YELLOW, 9.72),
    ]
    for num, label, body, accent, x in zones:
        rect(slide, x, 2.12, 2.55, 1.58, PANEL, accent, 4, True)
        text(slide, num, x + 0.2, 2.32, 0.42, 0.25, 9.5, accent, True, MONO)
        text(slide, label, x + 0.2, 2.72, 2.15, 0.28, 9.5, WHITE, True, MONO, tracking=18)
        text(slide, body, x + 0.2, 3.11, 2.06, 0.45, 10.3, MUTED, False, FONT)
    add_arrow(slide, 3.22, 2.68, "→", CYAN)
    add_arrow(slide, 6.22, 2.68, "→", BLUE)
    add_arrow(slide, 9.22, 2.68, "→", RED)

    # Request path and deny-by-default examples.
    text(slide, "REQUEST PATH", 0.75, 4.22, 1.55, 0.22, 8.5, CYAN, True, MONO, tracking=90)
    line(slide, 0.9, 4.82, 11.9, 4.82, GRID, 1.4)
    path = [("REQUEST", 1.0, CYAN), ("POLICY", 3.8, BLUE), ("WORKER", 6.6, RED), ("RESULT", 9.4, CYAN)]
    for label, x, accent in path:
        circle(slide, x, 4.63, 0.38, accent)
        text(slide, "•", x, 4.6, 0.38, 0.3, 15, BG, True, FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
        text(slide, label, x - 0.35, 5.12, 1.1, 0.2, 8, WHITE, True, MONO, PP_ALIGN.CENTER, tracking=35)

    rect(slide, 0.72, 5.72, 12.0, 0.86, "0D1422", GRID, 0, True)
    text(slide, "ALLOW", 1.0, 5.92, 0.65, 0.2, 8.5, CYAN, True, MONO, tracking=40)
    text(slide, "RED  →  LLM", 1.86, 5.9, 1.6, 0.22, 11, WHITE, True, MONO)
    text(slide, "BLOCK", 4.35, 5.92, 0.65, 0.2, 8.5, RED, True, MONO, tracking=40)
    text(slide, "RED  →  BLUE PRIVATE DATA", 5.2, 5.9, 2.85, 0.22, 11, WHITE, True, MONO)
    text(slide, "BLOCK", 8.45, 5.92, 0.65, 0.2, 8.5, RED, True, MONO, tracking=40)
    text(slide, "RED  →  HOST", 9.3, 5.9, 1.65, 0.22, 11, WHITE, True, MONO)
    text(slide, "default deny", 11.12, 5.92, 1.2, 0.22, 8.6, YELLOW, True, MONO, PP_ALIGN.RIGHT)
    return slide


def slide_05(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "quiet.jpg", tint=7)
    chrome(slide, 5, "BAYORA / PROOF")
    section_label(
        slide,
        "THE PROOF",
        "We don’t just say secure.\nWe show the decision.",
        "A reproducible validation run exercises policy enforcement, the isolated worker, and audit integrity in one view.",
        0.72,
        0.64,
        7.0,
    )

    # Large score block.
    rect(slide, 8.88, 0.92, 3.84, 1.52, "0E1E2A", CYAN, 0, True)
    text(slide, "6", 9.18, 1.08, 1.25, 0.75, 40, CYAN, True, DISPLAY)
    text(slide, "/ 6", 10.3, 1.3, 0.9, 0.35, 18, WHITE, True, DISPLAY)
    text(slide, "CONTROL CASES\nPASS IN FIXTURE MODE", 11.05, 1.16, 1.3, 0.58, 8, MUTED, True, MONO, tracking=35)
    text(slide, "application policy + fixture worker", 9.2, 2.11, 3.0, 0.16, 7.7, DIM, True, MONO, tracking=22)

    text(slide, "REPRODUCIBLE VALIDATION", 0.75, 2.72, 3.0, 0.22, 8.5, CYAN, True, MONO, tracking=90)
    rect(slide, 0.72, 3.1, 12.0, 3.18, PANEL, GRID, 0, True)
    rows = [
        ("RED  →  LLM", "ALLOW", "authorized evaluation flow", CYAN),
        ("RED  →  BLUE PRIVATE DATA", "BLOCK", "cross-role data boundary", RED),
        ("BLUE  →  RED PRIVATE DATA", "BLOCK", "cross-role data boundary", RED),
        ("RED  →  HOST", "BLOCK", "host resource boundary", RED),
        ("CROSS-SESSION DATA", "BLOCK", "session isolation boundary", RED),
        ("AUDIT TAMPERING", "DETECTED", "HMAC chain integrity", YELLOW),
    ]
    yy = 3.38
    for i, (flow, decision, control, accent) in enumerate(rows):
        if i:
            line(slide, 0.96, yy - 0.08, 12.45, yy - 0.08, GRID, 0.55)
        text(slide, f"{i+1:02d}", 0.98, yy, 0.35, 0.22, 8, DIM, True, MONO)
        text(slide, flow, 1.56, yy, 3.32, 0.22, 10, WHITE, True, MONO)
        pill(slide, decision, 5.18, yy - 0.02, 1.0 if decision != "DETECTED" else 1.22, accent, "101B2B", WHITE, 7.4)
        text(slide, control, 6.75, yy, 4.65, 0.22, 9.8, MUTED, False, FONT)
        text(slide, "PASS", 11.55, yy, 0.62, 0.22, 8, CYAN, True, MONO, PP_ALIGN.RIGHT, tracking=40)
        yy += 0.48
    text(slide, "The claim is bounded: application policy is proven; host-kernel sandboxing is not claimed.", 0.76, 6.54, 9.5, 0.2, 8.6, DIM, False, FONT, italic=True)
    return slide


def slide_06(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "quiet.jpg", tint=5)
    chrome(slide, 6, "BAYORA / CREDIBILITY")
    section_label(
        slide,
        "THE CREDIBILITY LAYER",
        "Designed for evidence — not theatre.",
        "Bayora makes the honest state visible: what ran, where it ran, which policy decided, and what a reviewer can verify.",
        0.72,
        0.64,
        8.2,
    )

    # Editable dashboard mockup, intentionally closer to a product screenshot than a generic UI card.
    rect(slide, 0.72, 2.02, 7.18, 4.55, "0A111C", GRID, 0, True)
    rect(slide, 0.72, 2.02, 7.18, 0.42, "101A2A", GRID, 0, True)
    circle(slide, 0.94, 2.17, 0.1, RED)
    circle(slide, 1.12, 2.17, 0.1, YELLOW)
    circle(slide, 1.3, 2.17, 0.1, CYAN)
    text(slide, "BAYORA  /  SECURITY EVALUATION", 1.62, 2.14, 3.4, 0.18, 7.4, MUTED, True, MONO, tracking=55)
    pill(slide, "ADMIN / ACTIVE", 6.36, 2.11, 1.1, CYAN, "0B2928", WHITE, 6.6)
    text(slide, "CONTROL PLANE", 1.02, 2.75, 2.3, 0.22, 8.2, CYAN, True, MONO, tracking=65)
    text(slide, "Provider", 1.02, 3.12, 0.85, 0.17, 7.2, DIM, True, MONO)
    text(slide, "FIXTURE", 1.02, 3.33, 1.08, 0.25, 14, WHITE, True, DISPLAY)
    text(slide, "Policy", 2.46, 3.12, 0.85, 0.17, 7.2, DIM, True, MONO)
    text(slide, "DEFAULT DENY", 2.46, 3.33, 1.55, 0.25, 14, WHITE, True, DISPLAY)
    text(slide, "Audit", 4.32, 3.12, 0.85, 0.17, 7.2, DIM, True, MONO)
    text(slide, "VERIFIED", 4.32, 3.33, 1.18, 0.25, 14, CYAN, True, DISPLAY)
    text(slide, "Role", 5.86, 3.12, 0.85, 0.17, 7.2, DIM, True, MONO)
    text(slide, "SEC ADMIN", 5.86, 3.33, 1.4, 0.25, 14, WHITE, True, DISPLAY)
    line(slide, 1.02, 3.88, 7.56, 3.88, GRID, 0.6)
    text(slide, "EVALUATION SESSIONS", 1.02, 4.1, 2.0, 0.18, 7.5, DIM, True, MONO, tracking=65)
    text(slide, "STATE", 6.54, 4.1, 0.75, 0.18, 7.5, DIM, True, MONO, PP_ALIGN.RIGHT, tracking=65)
    sessions = [("7f3c…b91a", "red-analyst", "ALLOW", CYAN), ("4a81…c220", "blue-review", "BLOCKED", RED), ("a01d…22fe", "auditor", "VERIFIED", CYAN)]
    sy = 4.5
    for sid, owner, state, accent in sessions:
        line(slide, 1.02, sy - 0.12, 7.56, sy - 0.12, GRID, 0.4)
        text(slide, sid, 1.02, sy, 1.45, 0.2, 9.2, WHITE, True, MONO)
        text(slide, owner, 2.7, sy, 1.55, 0.2, 8.8, MUTED, False, FONT)
        text(slide, "policy → worker → audit", 4.4, sy, 1.85, 0.2, 8.3, DIM, False, MONO)
        text(slide, state, 6.4, sy, 1.16, 0.2, 8, accent, True, MONO, PP_ALIGN.RIGHT)
        sy += 0.52
    text(slide, "TEST FIXTURE — NOT REAL MODEL EVIDENCE", 1.02, 6.06, 3.7, 0.2, 7.6, YELLOW, True, MONO, tracking=38)

    # Credibility reasons on the right.
    text(slide, "WHY JUDGES CAN TRUST THE DEMO", 8.42, 2.22, 3.8, 0.22, 8.5, CYAN, True, MONO, tracking=82)
    small_check(slide, 8.42, 2.77, "No silent fallback from real provider to fixture", CYAN)
    small_check(slide, 8.42, 3.48, "Raw prompts stay out of audit payloads", BLUE)
    small_check(slide, 8.42, 4.19, "Evidence type is visible and labeled", YELLOW)
    small_check(slide, 8.42, 4.9, "Every decision keeps tenant + session scope", PURPLE)
    rect(slide, 8.42, 5.72, 4.3, 0.84, "0E1C2C", CYAN, 0, True)
    text(slide, "SECURITY CLAIM", 8.7, 5.91, 1.26, 0.16, 7.2, DIM, True, MONO, tracking=55)
    text(slide, "→  mechanism  →  evidence", 10.06, 5.88, 2.33, 0.22, 9.8, WHITE, True, DISPLAY)
    return slide


def slide_07(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "quiet.jpg", tint=12)
    chrome(slide, 7, "BAYORA / WHY NOW")
    section_label(
        slide,
        "WHY NOW",
        "A wedge into the AI governance stack.",
        "As AI moves from prototype to production, the security workflow must move from prompt experiments to repeatable evidence.",
        0.72,
        0.64,
        8.6,
    )

    # Audience wedge on the left.
    text(slide, "FIRST USERS", 0.75, 2.08, 1.4, 0.22, 8.5, CYAN, True, MONO, tracking=90)
    audiences = [
        ("AI BUILDERS", "Ship safer features", CYAN),
        ("SECURITY RESEARCHERS", "Probe with guardrails", RED),
        ("ENTERPRISE SECOPS", "Review defensible evidence", BLUE),
    ]
    yy = 2.47
    for label, body, accent in audiences:
        rect(slide, 0.72, yy, 4.3, 0.72, PANEL, accent, 0, True)
        circle(slide, 0.98, yy + 0.22, 0.28, accent)
        text(slide, "•", 0.98, yy + 0.19, 0.28, 0.26, 12, BG, True, FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
        text(slide, label, 1.5, yy + 0.16, 2.2, 0.18, 8.8, WHITE, True, MONO, tracking=38)
        text(slide, body, 1.5, yy + 0.4, 2.85, 0.18, 9.5, MUTED, False, FONT)
        yy += 0.92

    # Commercial progression.
    text(slide, "THE WEDGE", 5.55, 2.08, 1.4, 0.22, 8.5, BLUE, True, MONO, tracking=90)
    wedge = [
        ("COMMUNITY CORE", "Repeatable probes + local validation", CYAN, 5.55, 2.47, 3.05),
        ("TEAM CONTROLS", "Shared findings + role-scoped review", BLUE, 5.55, 3.44, 3.55),
        ("ENTERPRISE EVIDENCE", "Identity, benchmarks, governance exports", PURPLE, 5.55, 4.41, 4.0),
    ]
    for label, body, accent, x, y, width in wedge:
        rect(slide, x, y, width, 0.72, PANEL, accent, 0, True)
        text(slide, label, x + 0.22, y + 0.14, width - 0.44, 0.18, 8.8, accent, True, MONO, tracking=35)
        text(slide, body, x + 0.22, y + 0.39, width - 0.44, 0.18, 9.4, WHITE, False, FONT)
    line(slide, 8.61, 3.19, 8.61, 3.42, CYAN, 1.2, 20)
    line(slide, 9.11, 4.16, 9.11, 4.39, BLUE, 1.2, 20)

    # Category context metric; explicitly labeled as external signal.
    rect(slide, 10.08, 2.12, 2.65, 3.0, "0F1728", GRID, 0, True)
    text(slide, "CATEGORY SIGNAL", 10.36, 2.43, 2.08, 0.18, 7.6, YELLOW, True, MONO, PP_ALIGN.CENTER, tracking=70)
    text(slide, "$308.3M", 10.31, 2.98, 2.16, 0.5, 26, WHITE, True, DISPLAY, PP_ALIGN.CENTER)
    text(slide, "AI governance market\nreference point", 10.34, 3.55, 2.1, 0.48, 9.3, MUTED, False, FONT, PP_ALIGN.CENTER)
    line(slide, 10.53, 4.26, 12.28, 4.26, GRID, 0.55)
    text(slide, "36% CAGR", 10.36, 4.5, 2.08, 0.34, 17, CYAN, True, DISPLAY, PP_ALIGN.CENTER)
    text(slide, "forecast category growth", 10.36, 4.9, 2.08, 0.18, 8.1, MUTED, False, FONT, PP_ALIGN.CENTER)
    text(slide, "External category context — not Bayora TAM.", 10.17, 5.38, 2.48, 0.28, 7.7, DIM, False, FONT, PP_ALIGN.CENTER)

    rect(slide, 0.72, 5.98, 12.0, 0.62, "0D1C2C", CYAN, 0, True)
    text(slide, "BUSINESS MODEL", 0.98, 6.18, 1.3, 0.18, 7.8, CYAN, True, MONO, tracking=60)
    text(slide, "Accessible core  →  paid team controls  →  enterprise evidence + assessment services", 2.52, 6.16, 8.85, 0.22, 10.2, WHITE, True, DISPLAY)
    text(slide, "land", 11.95, 6.17, 0.5, 0.18, 8, CYAN, True, MONO, PP_ALIGN.RIGHT)
    return slide


def slide_08(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "quiet.jpg", tint=10)
    chrome(slide, 8, "BAYORA / ROADMAP")
    section_label(
        slide,
        "THE ROADMAP",
        "Ship the wedge. Grow the moat.",
        "The MVP stays focused: authorized tests, clear boundaries, and an evidence trail that survives review.",
        0.72,
        0.64,
        8.5,
    )
    text(slide, "THE BUILD PATH", 0.75, 2.08, 1.6, 0.22, 8.5, CYAN, True, MONO, tracking=90)

    # Staircase timeline adds momentum without a generic three-column roadmap.
    milestones = [
        ("NOW", "PROVE THE WEDGE", "Repeatable authorized tests\nTenant + session boundaries\nSigned audit trail", CYAN, 0.72, 4.95, 3.25),
        ("NEXT", "EXPAND THE SIGNAL", "Provider adapters\nOWASP / ATLAS mappings\nReport + benchmark formats", BLUE, 4.27, 4.1, 3.25),
        ("SCALE", "OPERATE WITH TRUST", "SSO / MFA + collaboration\nMulti-org evidence\nEnterprise governance", PURPLE, 7.82, 3.25, 3.25),
    ]
    for phase, label, body, accent, x, y, width in milestones:
        rect(slide, x, y, width, 1.9, PANEL, accent, 0, True)
        rect(slide, x, y, width, 0.12, accent)
        pill(slide, phase, x + 0.24, y + 0.27, 0.74, accent, "101B2B", WHITE, 7.4)
        text(slide, label, x + 0.24, y + 0.72, width - 0.48, 0.25, 11.3, WHITE, True, DISPLAY)
        text(slide, body, x + 0.24, y + 1.12, width - 0.48, 0.58, 9.5, MUTED, False, FONT)
    add_arrow(slide, 3.84, 5.18, "→", CYAN)
    add_arrow(slide, 7.39, 4.32, "→", BLUE)

    # A crisp, defensible moat statement.
    rect(slide, 0.72, 2.55, 11.8, 0.52, "0C1B27", CYAN, 0, True)
    text(slide, "THE MOAT IS NOT A PROMPT LIST.", 1.02, 2.69, 3.25, 0.2, 9.3, CYAN, True, MONO, tracking=55)
    text(slide, "It is trustworthy evidence over time.", 4.46, 2.64, 4.5, 0.27, 16, WHITE, True, DISPLAY)
    text(slide, "reviewable", 10.92, 2.7, 1.25, 0.18, 8.4, CYAN, True, MONO, PP_ALIGN.RIGHT)

    # Bottom proof points.
    small_check(slide, 0.82, 6.45, "Focused MVP", CYAN)
    small_check(slide, 3.87, 6.45, "Visible limits", YELLOW)
    small_check(slide, 6.92, 6.45, "Composable adapters", BLUE)
    small_check(slide, 10.02, 6.45, "Enterprise path", PURPLE)
    return slide


def slide_09(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, "arena.jpg", tint=52)
    rect(slide, 0, 0, W, H, BG, transparency=44)
    # The centered overlay is intentionally quiet so the two masks read as a
    # red/blue visual metaphor rather than a stock-image background.
    rect(slide, 2.03, 0.75, 9.28, 5.85, "050912", transparency=16, radius=True)
    line(slide, 4.32, 1.15, 9.05, 1.15, CYAN, 1.7)
    text(slide, "THE TAKEAWAY", 4.32, 1.45, 4.7, 0.22, 8.8, CYAN, True, MONO, PP_ALIGN.CENTER, tracking=110)
    text(slide, "SECURE THE AI\nBEFORE USERS FIND THE EDGE.", 2.55, 2.04, 8.25, 1.38, 28, WHITE, True, DISPLAY, PP_ALIGN.CENTER, tracking=-16)
    rect(slide, 4.77, 3.76, 1.45, 0.08, RED)
    rect(slide, 6.29, 3.76, 1.45, 0.08, BLUE)
    text(slide, "Bayora makes responsible AI security\ntestable, repeatable, and reviewable.", 3.45, 4.2, 6.45, 0.64, 14, MUTED, False, FONT, PP_ALIGN.CENTER)
    pill(slide, "ATTACK THE MODEL", 4.02, 5.39, 1.74, RED, "2A121B")
    pill(slide, "PROVE THE BOUNDARY", 6.02, 5.39, 2.04, CYAN, "0D2826")
    text(slide, "BAYORA", 0.76, 6.82, 1.4, 0.2, 9, WHITE, True, MONO, tracking=90)
    text(slide, "[TEAM NAME]  /  HACK IN HILLS", 9.2, 6.82, 3.42, 0.2, 8, MUTED, True, MONO, PP_ALIGN.RIGHT, tracking=45)
    text(slide, "09", 12.0, 0.45, 0.6, 0.24, 8, CYAN, True, MONO, PP_ALIGN.RIGHT)
    return slide


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    # Remove the default title slide if a template version ever supplies one.
    while prs.slides:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]

    slide_01(prs)
    slide_02(prs)
    slide_03(prs)
    slide_04(prs)
    slide_05(prs)
    slide_06(prs)
    slide_07(prs)
    slide_08(prs)
    slide_09(prs)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
